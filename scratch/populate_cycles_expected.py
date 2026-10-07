import json
import os
import openpyxl
from collections import defaultdict

# 1. August Facilitator extraction
wb_aug = openpyxl.load_workbook('SS_ResponseDetail_Cluster Level_Grades 6-8_August.xlsx', data_only=True)
ws_aug = wb_aug['Facilitator']
rows_aug = list(ws_aug.iter_rows(values_only=True))
h_aug = rows_aug[0]
d_idx_a = h_aug.index('DistrictName')
b_idx_a = h_aug.index('BlockName')
q72_idx_a = h_aug.index('72')

aug_dist_exp = defaultdict(float)
aug_block_exp = defaultdict(lambda: defaultdict(float))

for r in rows_aug[1:]:
    d = str(r[d_idx_a]).strip()
    b = str(r[b_idx_a]).strip()
    val = r[q72_idx_a]
    if val is not None:
        try:
            v = float(str(val).strip().replace(',', ''))
            aug_dist_exp[d] += v
            aug_block_exp[d][b] += v
        except ValueError:
            pass

# 2. September Facilitator extraction
wb_sep = openpyxl.load_workbook(os.path.join('September_2026_Raw_Data', 'SS_ResponseDetail_Cluster Level_Grades 6-8_September.xlsx'), data_only=True)
ws_sep = wb_sep['Facilitator']
rows_sep = list(ws_sep.iter_rows(values_only=True))
h_sep = rows_sep[0]
d_idx_s = h_sep.index('DistrictName')
b_idx_s = h_sep.index('BlockName')
q72_idx_s = h_sep.index('72')

sep_dist_exp = defaultdict(float)
sep_block_exp = defaultdict(lambda: defaultdict(float))

for r in rows_sep[1:]:
    d = str(r[d_idx_s]).strip()
    b = str(r[b_idx_s]).strip()
    val = r[q72_idx_s]
    if val is not None:
        try:
            v = float(str(val).strip().replace(',', ''))
            sep_dist_exp[d] += v
            sep_block_exp[d][b] += v
        except ValueError:
            pass

dp = json.load(open('dataPackage.json', encoding='utf-8'))

# Update root districtSummary and blockSummary (Default = August)
for d in dp.get('districtSummary', []):
    dname = d['district']
    exp = aug_dist_exp.get(dname, 0)
    if exp == 0 and d.get('varg2Universe', 0) > 0:
        exp = round(d['varg2Universe'] * (68369 / 68427))
    d['expectedParticipants'] = exp
    d['targetCohort'] = exp
    if exp > 0 and 'attendees' in d:
        d['cohortTurnoutPct'] = round((d['attendees'] / exp) * 100, 1)

for b in dp.get('blockSummary', []):
    dname = b['district']
    bname = b['block']
    exp = aug_block_exp.get(dname, {}).get(bname, 0)
    if exp == 0 and b.get('varg2Universe', 0) > 0:
        exp = round(b['varg2Universe'] * (68369 / 68427))
    b['expectedParticipants'] = exp
    b['targetCohort'] = exp
    if exp > 0 and 'attendees' in b:
        b['cohortTurnoutPct'] = round((b['attendees'] / exp) * 100, 1)

# Update cycles.AUG
if 'cycles' in dp and 'AUG' in dp['cycles'] and isinstance(dp['cycles']['AUG'], dict):
    dp['cycles']['AUG']['varg2Metrics'] = {
        'totalUniverse': 68427,
        'targetCohort': 68369,
        'actualAttendees': 23785,
        'cadreSaturationPct': 34.8,
        'cohortTurnoutPct': round((23785 / 68369) * 100, 1)
    }
    for d in dp['cycles']['AUG'].get('districtSummary', []):
        dname = d['district']
        exp = aug_dist_exp.get(dname, 0)
        if exp == 0 and d.get('varg2Universe', 0) > 0:
            exp = round(d['varg2Universe'] * (68369 / 68427))
        d['expectedParticipants'] = exp
        d['targetCohort'] = exp
        if exp > 0 and 'attendees' in d:
            d['cohortTurnoutPct'] = round((d['attendees'] / exp) * 100, 1)
    for b in dp['cycles']['AUG'].get('blockSummary', []):
        dname = b['district']
        bname = b['block']
        exp = aug_block_exp.get(dname, {}).get(bname, 0)
        if exp == 0 and b.get('varg2Universe', 0) > 0:
            exp = round(b['varg2Universe'] * (68369 / 68427))
        b['expectedParticipants'] = exp
        b['targetCohort'] = exp
        if exp > 0 and 'attendees' in b:
            b['cohortTurnoutPct'] = round((b['attendees'] / exp) * 100, 1)

# Update cycles.SEP
if 'cycles' in dp and 'SEP' in dp['cycles'] and isinstance(dp['cycles']['SEP'], dict):
    dp['cycles']['SEP']['varg2Metrics'] = {
        'totalUniverse': 68427,
        'targetCohort': 67222,
        'actualAttendees': 23169,
        'cadreSaturationPct': 33.9,
        'cohortTurnoutPct': round((23169 / 67222) * 100, 1)
    }
    for d in dp['cycles']['SEP'].get('districtSummary', []):
        dname = d['district']
        exp = sep_dist_exp.get(dname, 0)
        if exp == 0 and d.get('varg2Universe', 0) > 0:
            exp = round(d['varg2Universe'] * (67222 / 68427))
        d['expectedParticipants'] = exp
        d['targetCohort'] = exp
        if exp > 0 and 'attendees' in d:
            d['cohortTurnoutPct'] = round((d['attendees'] / exp) * 100, 1)
    for b in dp['cycles']['SEP'].get('blockSummary', []):
        dname = b['district']
        bname = b['block']
        exp = sep_block_exp.get(dname, {}).get(bname, 0)
        if exp == 0 and b.get('varg2Universe', 0) > 0:
            exp = round(b['varg2Universe'] * (67222 / 68427))
        b['expectedParticipants'] = exp
        b['targetCohort'] = exp
        if exp > 0 and 'attendees' in b:
            b['cohortTurnoutPct'] = round((b['attendees'] / exp) * 100, 1)

# Update cycles.CONSOLIDATED
if 'cycles' in dp and 'CONSOLIDATED' in dp['cycles'] and isinstance(dp['cycles']['CONSOLIDATED'], dict):
    dp['cycles']['CONSOLIDATED']['varg2Metrics'] = {
        'totalUniverse': 68427,
        'targetCohort': 68369 + 67222,
        'actualAttendees': 23785 + 23169,
        'cadreSaturationPct': round(((23785 + 23169) / 68427) * 100, 1),
        'cohortTurnoutPct': round(((23785 + 23169) / (68369 + 67222)) * 100, 1)
    }
    for d in dp['cycles']['CONSOLIDATED'].get('districtSummary', []):
        dname = d['district']
        exp_a = aug_dist_exp.get(dname, 0) or round(d.get('varg2Universe', 700) * (68369 / 68427))
        exp_s = sep_dist_exp.get(dname, 0) or round(d.get('varg2Universe', 700) * (67222 / 68427))
        exp = exp_a + exp_s
        d['expectedParticipants'] = exp
        d['targetCohort'] = exp
        if exp > 0 and 'attendees' in d:
            d['cohortTurnoutPct'] = round((d['attendees'] / exp) * 100, 1)

with open('dataPackage.json', 'w', encoding='utf-8') as f:
    json.dump(dp, f, indent=2, ensure_ascii=False)

print('dataPackage.json successfully updated for all cycles!')
