import openpyxl
import os
import json

file_map = {
    'Cluster Level - August': 'SS_ResponseDetail_Cluster Level_Grades 6-8_August.xlsx',
    'Cluster Level - September': os.path.join('September_2026_Raw_Data', 'SS_ResponseDetail_Cluster Level_Grades 6-8_September.xlsx'),
    'District Level - August': 'SS_ResponseDetail_District Level_Grades 6-8_August.xlsx',
    'District Level - September': os.path.join('September_2026_Raw_Data', 'SS_ResponseDetail_District Level_Grades 6-8_September.xlsx')
}

target_phrase = 'अपेक्षित प्रतिभागियों'

results = []

for label, fpath in file_map.items():
    if not os.path.exists(fpath):
        continue
    wb = openpyxl.load_workbook(fpath, data_only=True)
    
    q_map = {}
    if 'Question Master' in wb.sheetnames:
        ws_q = wb['Question Master']
        for r in ws_q.iter_rows(values_only=True):
            if r and len(r) >= 4 and r[0] is not None:
                q_id = str(r[0])
                q_text = str(r[3])
                q_map[q_id] = q_text
    
    for sname in wb.sheetnames:
        if sname == 'Question Master':
            continue
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h) for h in rows[0]]
        for col_idx, h in enumerate(headers):
            is_match = False
            if target_phrase in h:
                is_match = True
            elif h in q_map and target_phrase in q_map[h]:
                is_match = True
            
            if is_match:
                vals = []
                for r in rows[1:]:
                    if col_idx < len(r) and r[col_idx] is not None:
                        try:
                            v = float(str(r[col_idx]).strip().replace(',', ''))
                            vals.append(v)
                        except ValueError:
                            pass
                col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
                total_sum = sum(vals)
                results.append({
                    'workbook': label,
                    'sheet': sname,
                    'col_letter': col_letter,
                    'col_index': col_idx + 1,
                    'col_header': h,
                    'question_text': q_map.get(h, h),
                    'rows': len(vals),
                    'sum': total_sum,
                    'average': (total_sum / len(vals)) if vals else 0
                })

with open('scratch/expected_teachers_exact_sums.json', 'w', encoding='utf-8') as f:
    json.dump(results, f, indent=2, ensure_ascii=False)

print('WROTE JSON TO scratch/expected_teachers_exact_sums.json')
